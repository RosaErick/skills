#!/usr/bin/env python3
"""Extract slide text, geometry, notes and images. This is not a visual renderer."""
import argparse
import json
from pathlib import Path


def extract(source, output):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    presentation = Presentation(source)
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    images = output / 'images'
    images.mkdir(exist_ok=True)
    result = {'source': Path(source).name, 'width_emu': presentation.slide_width,
              'height_emu': presentation.slide_height, 'slides': [],
              'limitations': 'Content extraction only; reconstruct layouts, charts, media and animations as needed.'}
    def shapes(items, slide_number, prefix=''):
        records = []
        for index, shape in enumerate(items, 1):
            identifier = f'{prefix}{index}'
            item = {'name': shape.name, 'type': str(shape.shape_type), 'left': shape.left,
                    'top': shape.top, 'width': shape.width, 'height': shape.height}
            if shape.has_text_frame:
                item['text'] = shape.text_frame.text
            if shape.has_table:
                item['table'] = [[cell.text for cell in row.cells] for row in shape.table.rows]
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                item['children'] = shapes(shape.shapes, slide_number, identifier + '-')
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                asset = images / f'slide-{slide_number}-shape-{identifier}.{shape.image.ext}'
                asset.write_bytes(shape.image.blob)
                item['image'] = str(asset.relative_to(output))
            if shape.has_chart:
                item['needs_reconstruction'] = 'chart'
            records.append(item)
        return records
    for number, slide in enumerate(presentation.slides, 1):
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide and slide.notes_slide.notes_text_frame else ''
        result['slides'].append({'number': number, 'notes': notes, 'shapes': shapes(slide.shapes, number)})
    destination = output / 'slides.json'
    destination.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding='utf-8')
    return destination


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('presentation', type=Path)
    parser.add_argument('output_directory', type=Path)
    args = parser.parse_args()
    try:
        print(extract(args.presentation, args.output_directory))
    except ImportError:
        parser.exit(2, 'Missing python-pptx; install the adjacent requirements.txt in an isolated environment.\n')
    except (OSError, ValueError) as error:
        parser.exit(1, f'Extraction failed: {error}\n')


if __name__ == '__main__':
    main()
