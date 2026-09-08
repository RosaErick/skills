import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[2]


def module(name, relative):
    spec = importlib.util.spec_from_file_location(name, ROOT / relative)
    loaded = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(loaded)
    return loaded


api = module('api_validator', 'backend/api-documentation-master/scripts/api_validator.py')
security = module('security_scan', 'security/vulnerability-scanner/scripts/security_scan.py')
lint = module('lint_runner', 'quality/lint-and-validate/scripts/lint_runner.py')
geo = module('geo_checker', 'writing/geo-fundamentals/scripts/geo_checker.py')
extractor = module('extract_pptx', 'frontend/frontend-slides/scripts/extract_pptx.py')


class Fixture(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='skill checks ')
        self.root = Path(self.temp.name)
        self.addCleanup(self.temp.cleanup)

    def write(self, name, content):
        path = self.root / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(content)
        return path


@unittest.skipUnless(importlib.util.find_spec('openapi_spec_validator'), 'install OpenAPI validator requirements')
class OpenAPITests(Fixture):
    def valid(self):
        return {'openapi': '3.1.0', 'info': {'title': 'Fixture', 'version': '1.0'},
                'paths': {'/items': {'get': {'responses': {'200': {'description': 'OK'}}}}}}

    def test_valid_json_and_yaml(self):
        import yaml
        spec = self.valid()
        for name, content in [('openapi.json', json.dumps(spec)), ('openapi.yml', yaml.safe_dump(spec))]:
            with self.subTest(name=name):
                self.assertEqual(api.check_openapi_spec(self.write(name, content))['status'], 'valid')

    def test_comments_and_malformed_schema_cannot_pass(self):
        fake = '# openapi: fake\n# paths: fake\n# components: fake\nthis is not an OpenAPI document\n'
        self.assertEqual(api.check_openapi_spec(self.write('fake.yaml', fake))['status'], 'invalid')
        spec = self.valid()
        spec['paths']['/items']['get']['responses'] = 'wrong type'
        self.assertEqual(api.check_openapi_spec(self.write('bad.json', json.dumps(spec)))['status'], 'invalid')

    def test_local_reference_and_unresolved_reference(self):
        spec = self.valid()
        spec['paths']['/items']['get']['responses']['200'] = {'$ref': './response.yaml'}
        self.write('response.yaml', 'description: OK\n')
        path = self.write('openapi.json', json.dumps(spec))
        self.assertEqual(api.check_openapi_spec(path)['status'], 'valid')
        (self.root / 'response.yaml').unlink()
        self.assertEqual(api.check_openapi_spec(path)['status'], 'invalid')

    def test_remote_and_escaping_references_do_not_fetch(self):
        spec = self.valid()
        for reference in ['http://127.0.0.1:9/private', '../outside.yaml']:
            with self.subTest(reference=reference):
                spec['paths']['/items']['get']['responses']['200'] = {'$ref': reference}
                path = self.write('openapi.json', json.dumps(spec))
                with patch('urllib.request.urlopen', side_effect=AssertionError('network is forbidden')) as request:
                    result = api.check_openapi_spec(path)
                self.assertEqual(result['status'], 'unverified', result)
                request.assert_not_called()

    def test_discovery_deduplicates_and_does_not_silently_limit(self):
        for index in range(18):
            self.write(f'{index}/api.openapi.json', json.dumps(self.valid()))
        self.write('node_modules/ignore.openapi.json', '{}')
        self.assertEqual(len(api.find_api_files(self.root)), 18)

    def test_missing_input_is_not_success(self):
        process = subprocess.run([sys.executable, str(ROOT / 'backend/api-documentation-master/scripts/api_validator.py'), str(self.root)], capture_output=True, text=True)
        self.assertEqual(process.returncode, 2)
        self.assertEqual(json.loads(process.stdout)['checked_files'], 0)


class DependencyTests(Fixture):
    def package(self, manager='pnpm@10.0.0', lock='pnpm-lock.yaml'):
        self.write('package.json', json.dumps({'packageManager': manager}))
        self.write(lock, '')

    def test_pnpm_does_not_require_npm_or_yarn_locks_or_contact_registry_by_default(self):
        self.package()
        with patch.object(security.subprocess, 'run') as run:
            result = security.scan_dependencies(str(self.root))
        self.assertEqual(result['manager'], 'pnpm')
        self.assertTrue(result['lockfile_found'])
        self.assertEqual(result['findings'], [])
        self.assertEqual(result['audit_status'], 'not_requested')
        run.assert_not_called()

    def test_matching_manager_audit_and_all_severities_reported(self):
        self.package()
        payload = {'metadata': {'vulnerabilities': {'low': 1, 'moderate': 2, 'high': 0, 'critical': 0, 'total': 3}}}
        with patch.object(security.subprocess, 'run', return_value=subprocess.CompletedProcess([], 1, json.dumps(payload), '')) as run:
            result = security.scan_dependencies(str(self.root), audit=True)
        self.assertEqual(run.call_args.args[0], ['pnpm', 'audit', '--json'])
        self.assertEqual(result['audit_status'], 'completed')
        self.assertEqual({finding['severity'] for finding in result['findings']}, {'low', 'moderate'})

    def test_invalid_unavailable_and_registry_error_audits_are_unverified(self):
        self.package('npm@11.0.0', 'package-lock.json')
        cases = [subprocess.CompletedProcess([], 1, 'not json', ''),
                 subprocess.CompletedProcess([], 0, '{"error":"offline"}', ''),
                 subprocess.CompletedProcess([], 1, '{"metadata":{"vulnerabilities":{"high":0}}}', '')]
        for output in cases:
            with patch.object(security.subprocess, 'run', return_value=output):
                result = security.scan_dependencies(str(self.root), audit=True)
            self.assertEqual(result['audit_status'], 'unverified')
        with patch.object(security.subprocess, 'run', side_effect=FileNotFoundError):
            self.assertEqual(security.scan_dependencies(str(self.root), audit=True)['audit_status'], 'unverified')

    def test_conflicting_lockfiles_do_not_pick_a_manager_silently(self):
        self.package()
        self.write('yarn.lock', '')
        with patch.object(security.subprocess, 'run') as run:
            result = security.scan_dependencies(str(self.root), audit=True)
        self.assertIn('Conflicting', result['notes'][0])
        run.assert_not_called()

    def test_classic_and_modern_yarn_use_their_supported_audit_commands(self):
        for version, response, expected in [
            ('1.22.0', '{"type":"auditSummary","data":{"vulnerabilities":{"high":0}}}', ['yarn', 'audit', '--json']),
            ('4.0.0', '{"metadata":{"vulnerabilities":{"high":0}}}', ['yarn', 'npm', 'audit', '--json', '--all', '--recursive'])]:
            self.package('yarn@' + version, 'yarn.lock')
            with patch.object(security.subprocess, 'run', return_value=subprocess.CompletedProcess([], 0, response, '')) as run:
                result = security.scan_dependencies(str(self.root), audit=True)
            self.assertEqual(run.call_args.args[0], expected)
            self.assertEqual(result['audit_status'], 'completed')

    def test_regex_matches_are_redacted_candidates(self):
        secret = 'fictional-secret-value-for-testing'
        self.write('app.js', f'const token = "{secret}"; eval(input);')
        report = security.run_full_scan(str(self.root), 'secrets')
        self.assertNotIn(secret, json.dumps(report))
        self.assertTrue(report['scans']['secrets']['findings'])
        self.assertTrue(all('not confirmed' in item['evidence'] for item in report['scans']['secrets']['findings']))

    def test_environment_files_are_scanned_but_public_ssh_keys_are_not_secrets(self):
        self.write('.env.local', 'TOKEN="fictional-test-token-value"')
        self.write('public_key.js', 'const publicKey = "ssh-rsa AAAAB3NzaC1yc2EAAAADAQABAAABAQ";')
        report = security.run_full_scan(str(self.root), 'secrets')
        findings = report['scans']['secrets']['findings']
        self.assertTrue(any(item['file'] == '.env.local' for item in findings))
        self.assertFalse(any(item['file'] == 'public_key.js' for item in findings))

    def test_pattern_findings_never_echo_source_secrets(self):
        secret = 'fixture-sensitive-string-do-not-echo'
        self.write('app.js', f'eval(input); const password = "{secret}";')
        report = security.run_full_scan(str(self.root), 'patterns')
        self.assertTrue(report['scans']['code_patterns']['findings'])
        self.assertNotIn(secret, json.dumps(report))


class LintTests(Fixture):
    def test_detects_configured_scripts_and_manager_without_inventing_commands(self):
        self.write('package.json', json.dumps({'packageManager': 'pnpm@10', 'scripts': {'lint': 'eslint .', 'typecheck': 'tsc --noEmit'}, 'devDependencies': {'typescript': '5'}}))
        detected = lint.detect_project_type(self.root)
        self.assertEqual([item['cmd'] for item in detected['linters']], [['pnpm', 'run', 'lint'], ['pnpm', 'run', 'typecheck']])

    def test_empty_python_tool_configuration_does_not_invent_mypy(self):
        self.write('pyproject.toml', '[project]\nname="fixture"\n')
        self.assertEqual(lint.detect_project_type(self.root)['linters'], [])

    def test_no_checks_exits_unverified(self):
        result = subprocess.run([sys.executable, str(ROOT / 'quality/lint-and-validate/scripts/lint_runner.py'), str(self.root)], capture_output=True, text=True)
        self.assertEqual(result.returncode, 2)
        self.assertIsNone(json.loads(result.stdout)['passed'])

    def test_failed_and_missing_checks_are_distinct(self):
        missing = lint.run_linter({'name': 'missing', 'cmd': ['no-such-skill-test-tool-47']}, self.root)
        self.assertEqual(missing['status'], 'unverified')
        failure = lint.run_linter({'name': 'fixture', 'cmd': [sys.executable, '-c', 'raise SystemExit(3)']}, self.root)
        self.assertEqual(failure['status'], 'failed')


class ContentTests(Fixture):
    def test_geo_observations_do_not_claim_ranking_or_originality(self):
        path = self.write('page.html', '<h1>Fixture</h1><p>100% invented statistics</p>')
        result = geo.check_page(path)
        self.assertEqual(result['status'], 'observed')
        self.assertNotIn('score', result)
        self.assertNotIn('passed', result)
        self.assertEqual(result['signals']['h1_tags'], 1)

    @unittest.skipUnless(importlib.util.find_spec('pptx'), 'install slide extractor requirements')
    def test_pptx_preserves_slide_order_notes_table_and_embedded_image(self):
        from pptx import Presentation
        from pptx.util import Inches
        import base64
        image = self.root / 'pixel.png'
        image.write_bytes(base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII='))
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = 'First'
        slide.notes_slide.notes_text_frame.text = 'Keep this speaker note'
        slide.shapes.add_picture(str(image), Inches(1), Inches(1))
        table = slide.shapes.add_table(1, 1, Inches(2), Inches(2), Inches(1), Inches(1)).table
        table.cell(0, 0).text = 'Table content'
        slide2 = deck.slides.add_slide(deck.slide_layouts[5])
        slide2.shapes.title.text = 'Second'
        source = self.root / 'source.pptx'
        deck.save(source)
        output = self.root / 'output'
        result = json.loads(extractor.extract(source, output).read_text())
        self.assertEqual([slide['number'] for slide in result['slides']], [1, 2])
        self.assertEqual(result['slides'][0]['notes'], 'Keep this speaker note')
        self.assertEqual(result['slides'][1]['shapes'][0]['text'], 'Second')
        self.assertIn([['Table content']], [shape.get('table') for shape in result['slides'][0]['shapes']])
        asset = next(shape['image'] for shape in result['slides'][0]['shapes'] if 'image' in shape)
        self.assertEqual((output / asset).read_bytes(), image.read_bytes())


if __name__ == '__main__':
    unittest.main()
